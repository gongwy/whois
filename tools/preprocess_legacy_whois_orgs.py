#!/usr/bin/env python3
"""First, reproducible legacy-WHOIS organisation pre-processing stage.

One command parses old IP/domain objects, reports name/association conflicts,
and resolves only multi-name conn-i values by legal-name existence in the
approved order ICP -> domain registration -> non-reserved IPFP.  It does not
materialise new WHOIS objects or use filing IP/domain ownership joins.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path
import subprocess
from typing import Any

from openpyxl import load_workbook

from migration_special_rules import (
    legacy_resource_exclusion_audit,
    legacy_resource_exclusion_for,
)
from prepare_joint_rsdb_stage1 import (
    DOMAIN_REGISTRATION_REQUIRED_HEADERS,
    IP_ALLOCATION_SHEET,
    ICP_SHEET,
    canonical_domain,
    first,
    legacy_names,
    meaningful,
    normalized,
    normalized_value,
    parse_ip_range,
    parse_rpsl_file,
    read_sheet,
    require_headers,
    text,
)


SOURCES = (
    ("inetnum", "CERNIC.db.in.txt", "inetnum:"),
    ("inet6num", "CERNIC.db.i6.txt", "*i6:"),
)


def distinct(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def joined(values: list[str]) -> str:
    return "\n".join(distinct(values))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_ref(obj: dict[str, Any]) -> str:
    return f"{obj['source_file']}:{obj['source_line']}"


def read_legacy(rsdb_dir: Path) -> dict[str, Any]:
    """Use precisely the prior conn-i and domain association rules."""
    names: dict[str, dict[str, str]] = defaultdict(dict)
    resources: dict[str, list[dict[str, Any]]] = defaultdict(list)
    netname_index: dict[str, set[str]] = defaultdict(set)
    name_index: dict[str, set[str]] = defaultdict(set)
    exclusions: list[dict[str, str]] = []
    ignored_no_conn_i: list[dict[str, str]] = []
    ignored_no_name: list[dict[str, str]] = []

    for kind, filename, marker in SOURCES:
        path = rsdb_dir / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        for obj in parse_rpsl_file(path, marker, kind):
            rule = legacy_resource_exclusion_for(kind, obj["raw_key"], obj["fields"])
            if rule:
                exclusions.append(legacy_resource_exclusion_audit(
                    rule, source_file=obj["source_file"], source_line=obj["source_line"]
                ))
                continue
            fields = obj["fields"]
            conn_i = first(fields, "conn-i").upper()
            interval = parse_ip_range(obj["raw_key"])
            if not conn_i or not interval:
                ignored_no_conn_i.append({
                    "object_type": kind, "resource": obj["raw_key"],
                    "conn_i": conn_i, "source_ref": source_ref(obj),
                    "reason": "缺少 conn-i 或 IP 区间无法解析",
                })
                continue
            orgname, english_name = legacy_names(fields)
            key = normalized(orgname)
            resources[conn_i].append({
                "object_type": kind,
                "resource": obj["raw_key"],
                "conn_i": conn_i,
                "netname": first(fields, "netname"),
                "orgname": orgname,
                "english_name": english_name,
                "source_ref": source_ref(obj),
            })
            netname = normalized_value(first(fields, "netname"))
            if netname:
                netname_index[netname].add(conn_i)
            if not key:
                ignored_no_name.append({
                    "object_type": kind, "resource": obj["raw_key"],
                    "conn_i": conn_i, "source_ref": source_ref(obj),
                    "reason": "旧 IP 对象少于两条 descr 或第一条机构名称为空",
                })
                continue
            names[conn_i].setdefault(key, orgname)
            name_index[key].add(conn_i)

    orgs = [
        {
            "conn_i": conn_i,
            "orgname_count": len(candidate_map),
            "orgnames": joined(list(candidate_map.values())),
            "resource_count": len(resources[conn_i]),
            "status": "同一conn-i多名称" if len(candidate_map) > 1 else "唯一名称",
        }
        for conn_i, candidate_map in sorted(names.items())
    ]
    org_conflicts = [row for row in orgs if row["orgname_count"] > 1]
    org_conflict_resources = [
        resource
        for row in org_conflicts
        for resource in resources[row["conn_i"]]
    ]

    domain_path = rsdb_dir / "CERNIC.db.dn.txt"
    if not domain_path.is_file():
        raise FileNotFoundError(domain_path)
    domains: list[dict[str, Any]] = []
    for obj in parse_rpsl_file(domain_path, "domain:", "domain"):
        fields = obj["fields"]
        netname = first(fields, "netname")
        c_orgname = first(fields, "c-orgname")
        net_candidates = sorted(netname_index.get(normalized_value(netname), set())) if netname else []
        name_candidates = sorted(name_index.get(normalized(c_orgname), set())) if c_orgname else []
        selected = net_candidates if net_candidates else name_candidates
        route = "netname" if net_candidates else "c-orgname" if name_candidates else "无"
        status = "唯一匹配" if len(selected) == 1 else "多候选冲突" if selected else "零匹配冲突"
        domains.append({
            "domain": obj["raw_key"],
            "canonical_domain": canonical_domain(obj["raw_key"]) or "",
            "netname": netname,
            "c_orgname": c_orgname,
            "netname_candidates": joined(net_candidates),
            "c_orgname_candidates": joined(name_candidates),
            "matching_route": route,
            "matched_conn_i": selected[0] if len(selected) == 1 else "",
            "candidate_conn_i": joined(selected),
            "candidate_count": len(selected),
            "status": status,
            "source_ref": source_ref(obj),
        })

    return {
        "organisations": orgs,
        "org_conflicts": org_conflicts,
        "org_conflict_resources": org_conflict_resources,
        "domains": domains,
        "domain_multiple": [row for row in domains if row["status"] == "多候选冲突"],
        "domain_zero": [row for row in domains if row["status"] == "零匹配冲突"],
        "excluded_legacy_resources": exclusions,
        "ignored_ip_without_conn_i_or_range": ignored_no_conn_i,
        "ignored_ip_without_orgname": ignored_no_name,
    }


def domain_registration_sheet(path: Path) -> str:
    book = load_workbook(path, read_only=True, data_only=True)
    matching = []
    try:
        for sheet in book.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = {text(value) for value in next(iterator)}
            except StopIteration:
                continue
            if DOMAIN_REGISTRATION_REQUIRED_HEADERS.issubset(headers):
                matching.append(sheet.title)
    finally:
        book.close()
    if len(matching) != 1:
        raise ValueError(f"域名备案文件必须且只能有一个符合字段的工作表: {matching}")
    return matching[0]


def read_filing_exclusions(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(f"缺少已确认的备案排除规则: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        required = {"source_kind", "organisation_name", "cert_no", "action", "reason"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError(f"备案排除规则字段不完整: {path}")
        return [{key: text(value) for key, value in row.items()} for row in reader]


def matching_filing_exclusion(
    rules: list[dict[str, str]], source_kind: str, name: str, cert_no: str,
) -> dict[str, str] | None:
    for rule in rules:
        if normalized(rule["source_kind"]) != normalized(source_kind):
            continue
        if normalized(rule["organisation_name"]) != normalized(name):
            continue
        rule_cert = normalized(meaningful(rule["cert_no"]))
        if rule_cert and rule_cert != normalized(meaningful(cert_no)):
            continue
        return rule
    return None


def load_legal_name_indexes(
    icp_path: Path, domain_path: Path, ip_path: Path,
    filing_exclusions: list[dict[str, str]],
) -> tuple[dict[str, dict[str, list[dict[str, Any]]]], list[dict[str, str]], int]:
    sources: dict[str, dict[str, list[dict[str, Any]]]] = {}
    excluded_rows: list[dict[str, str]] = []
    reserved_ipfp_rows = 0

    icp_rows = read_sheet(icp_path, ICP_SHEET)
    require_headers(icp_rows, {"主体单位名称", "主办单位有效证件号码", "域名"}, ICP_SHEET)
    icp_index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_no, row in icp_rows:
        name = meaningful(row["主体单位名称"])
        rule = matching_filing_exclusion(filing_exclusions, "icp", name, row["主办单位有效证件号码"])
        if rule:
            excluded_rows.append({"source_kind": "ICP", "source_ref": f"{icp_path.name}/{ICP_SHEET}!{row_no}", "legal_name": name, "reason": rule["reason"]})
            continue
        if normalized(name):
            icp_index[normalized(name)].append({
                "legal_name": name, "record_id": f"ICP-{row_no}",
                "source_ref": f"{icp_path.name}/{ICP_SHEET}!{row_no}",
                "domain": row["域名"],
                "certificate_type": row.get("主办单位有效证件类型", ""),
                "certificate_number": row["主办单位有效证件号码"],
            })
    sources["ICP"] = icp_index

    sheet_name = domain_registration_sheet(domain_path)
    domain_rows = read_sheet(domain_path, sheet_name)
    require_headers(domain_rows, DOMAIN_REGISTRATION_REQUIRED_HEADERS, sheet_name)
    domain_index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_no, row in domain_rows:
        name = meaningful(row["registrantName"]) or meaningful(row["registrantOrg"])
        rule = matching_filing_exclusion(filing_exclusions, "domain-registration", name, row["registrantCertNo"])
        if rule:
            excluded_rows.append({"source_kind": "DOMAIN_REGISTRY", "source_ref": f"{domain_path.name}/{sheet_name}!{row_no}", "legal_name": name, "reason": rule["reason"]})
            continue
        if normalized(name):
            domain_index[normalized(name)].append({
                "legal_name": name, "record_id": f"DOMAINREG-{row_no}",
                "source_ref": f"{domain_path.name}/{sheet_name}!{row_no}",
                "domain": joined([row["name"], row["punycode"]]),
                "certificate_type": row["registrantCertType"],
                "certificate_number": row["registrantCertNo"],
            })
    sources["DOMAIN_REGISTRY"] = domain_index

    ip_rows = read_sheet(ip_path, IP_ALLOCATION_SHEET)
    require_headers(ip_rows, {"ID", "分配方式", "使用单位名称", "使用单位证件号码", "起始IP", "终止IP"}, IP_ALLOCATION_SHEET)
    ip_index: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row_no, row in ip_rows:
        if normalized(row["分配方式"]) == normalized("预留"):
            reserved_ipfp_rows += 1
            continue
        name = meaningful(row["使用单位名称"])
        rule = matching_filing_exclusion(filing_exclusions, "ip-allocation", name, row["使用单位证件号码"])
        if rule:
            excluded_rows.append({"source_kind": "IPFP", "source_ref": f"{ip_path.name}/{IP_ALLOCATION_SHEET}!{row_no}", "legal_name": name, "reason": rule["reason"]})
            continue
        if normalized(name):
            ip_index[normalized(name)].append({
                "legal_name": name, "record_id": f"IP-{row['ID'] or row_no}",
                "source_ref": f"{ip_path.name}/{IP_ALLOCATION_SHEET}!{row_no}",
                "domain": "", "ip_range_audit_only": f"{row['起始IP']} - {row['终止IP']}",
                "certificate_type": row.get("使用单位证件类型", ""),
                "certificate_number": row["使用单位证件号码"],
            })
    sources["IPFP"] = ip_index
    return sources, excluded_rows, reserved_ipfp_rows


def resolve_one_stage(
    previous: list[dict[str, Any]],
    index: dict[str, list[dict[str, Any]]],
    source_kind: str,
) -> dict[str, Any]:
    """Only previous unresolved rows enter; one name hit selects old WHOIS text."""
    results: list[dict[str, Any]] = []
    candidates: list[dict[str, Any]] = []
    evidence: list[dict[str, Any]] = []
    for item in previous:
        if item.get("selected_orgname"):
            continue
        conn_i = item["conn_i"]
        names = item["orgnames"].splitlines()
        matched: list[str] = []
        for name in names:
            records = index.get(normalized(name), [])
            if records:
                matched.append(name)
            candidates.append({
                "conn_i": conn_i, "old_orgname": name,
                "matched": "是" if records else "否", "filing_record_count": len(records),
                "filing_names": joined([record["legal_name"] for record in records]),
                "filing_sources": joined([record["source_ref"] for record in records]),
            })
            for record in records:
                evidence.append({
                    "conn_i": conn_i, "old_orgname": name,
                    "filing_name": record["legal_name"],
                    "record_id": record["record_id"],
                    "source_ref": record["source_ref"],
                    "domain": record.get("domain", ""),
                    "ip_range_audit_only": record.get("ip_range_audit_only", ""),
                    "certificate_type": record.get("certificate_type", ""),
                    "certificate_number": record.get("certificate_number", ""),
                })
        if len(matched) == 1:
            status = "唯一命中"
            selected = matched[0]
        elif len(matched) > 1:
            status = "多个名称命中"
            selected = ""
        else:
            status = "未命中"
            selected = ""
        results.append({
            "conn_i": conn_i,
            "source_kind": source_kind,
            "prior_status": item["status"],
            "old_orgname_count": len(names),
            "old_orgnames": item["orgnames"],
            "matched_name_count": len(matched),
            "matched_old_orgnames": joined(matched),
            "selected_orgname": selected,
            "status": status,
            "resource_count": item["resource_count"],
        })
    return {
        "source_kind": source_kind,
        "results": results,
        "candidates": candidates,
        "evidence": evidence,
        "summary": {
            "input_conn_i": len(results),
            "unique": sum(row["status"] == "唯一命中" for row in results),
            "multiple": sum(row["status"] == "多个名称命中" for row in results),
            "none": sum(row["status"] == "未命中" for row in results),
            "candidate_rows": len(candidates),
            "evidence_rows": len(evidence),
        },
    }


def build_report(
    legacy: dict[str, Any],
    indexes: dict[str, dict[str, list[dict[str, Any]]]],
) -> dict[str, Any]:
    unresolved = [{**row, "status": "旧WHOIS多名称"} for row in legacy["org_conflicts"]]
    stages = []
    decisions: dict[str, dict[str, str]] = {}
    for source_kind in ("ICP", "DOMAIN_REGISTRY", "IPFP"):
        stage = resolve_one_stage(unresolved, indexes[source_kind], source_kind)
        stages.append(stage)
        for row in stage["results"]:
            if row["selected_orgname"]:
                decisions[row["conn_i"]] = {
                    "conn_i": row["conn_i"],
                    "selected_orgname": row["selected_orgname"],
                    "selected_at_stage": source_kind,
                }
        stage_by_id = {row["conn_i"]: row for row in stage["results"]}
        unresolved = [
            {**row, "status": stage_by_id[row["conn_i"]]["status"]}
            for row in unresolved if not stage_by_id[row["conn_i"]]["selected_orgname"]
        ]
    summary = {
        "orgid_total": len(legacy["organisations"]),
        "multi_name_conn_i": len(legacy["org_conflicts"]),
        "domain_total": len(legacy["domains"]),
        "domain_unique": sum(row["status"] == "唯一匹配" for row in legacy["domains"]),
        "domain_multiple": len(legacy["domain_multiple"]),
        "domain_zero": len(legacy["domain_zero"]),
        "excluded_legacy_resources": len(legacy["excluded_legacy_resources"]),
        "ignored_ip_without_conn_i_or_range": len(legacy["ignored_ip_without_conn_i_or_range"]),
        "ignored_ip_without_orgname": len(legacy["ignored_ip_without_orgname"]),
        "selected_multi_name_conn_i": len(decisions),
        "remaining_multi_name_conn_i": len(unresolved),
    }
    return {
        "stage": "legacy-whois-org-preprocessing",
        "status": "completed_with_unresolved_conflicts" if unresolved or legacy["domain_multiple"] or legacy["domain_zero"] else "passed",
        "rules": {
            "orgid": "旧inetnum/inet6num的conn-i；机构名称候选为第一条descr（至少两条descr）",
            "domain_link": "netname先查旧IP netname；零候选时用域名c-orgname查旧IP第一条descr；不得使用备案倒推旧conn-i",
            "name_stage_order": ["ICP", "DOMAIN_REGISTRY", "IPFP"],
            "name_match": "仅比对法人名称是否存在；规范化只用于索引，确定时保留旧WHOIS原始机构名称",
            "ipfp_reserved": "完全排除分配方式为预留的IP备案行",
            "resource_preservation": "旧domain与IP资源主键只保留源值；本阶段不物化任何新WHOIS对象",
        },
        "summary": summary,
        "legacy": legacy,
        "name_stages": stages,
        "selected_multi_name_orgnames": sorted(decisions.values(), key=lambda row: row["conn_i"]),
        "unresolved_multi_name_orgs": unresolved,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rsdb-dir", type=Path, required=True)
    parser.add_argument("--icp-workbook", type=Path, required=True)
    parser.add_argument("--domain-workbook", type=Path, required=True)
    parser.add_argument("--ip-workbook", type=Path, required=True)
    parser.add_argument("--filing-exclusions", type=Path, help="已确认的备案排除规则TSV；默认从IP备案文件同目录读取")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--no-xlsx", action="store_true", help="仅输出机器可读JSON")
    args = parser.parse_args()
    legacy = read_legacy(args.rsdb_dir)
    exclusion_path = args.filing_exclusions or args.ip_workbook.parent / "迁移排除记录.tsv"
    indexes, excluded_filing_rows, reserved_ipfp_rows = load_legal_name_indexes(
        args.icp_workbook, args.domain_workbook, args.ip_workbook,
        read_filing_exclusions(exclusion_path),
    )
    report = build_report(legacy, indexes)
    report["excluded_filing_rows"] = excluded_filing_rows
    report["summary"]["excluded_filing_rows"] = len(excluded_filing_rows)
    report["summary"]["excluded_reserved_ipfp_rows"] = reserved_ipfp_rows
    source_paths = [
        args.rsdb_dir / "CERNIC.db.in.txt",
        args.rsdb_dir / "CERNIC.db.i6.txt",
        args.rsdb_dir / "CERNIC.db.dn.txt",
        args.icp_workbook, args.domain_workbook, args.ip_workbook, exclusion_path,
    ]
    report["sources"] = [
        {"filename": path.name, "sha256": sha256(path), "bytes": path.stat().st_size}
        for path in source_paths
    ]
    if args.output_dir.exists() and any(args.output_dir.iterdir()):
        raise FileExistsError(f"输出目录已有文件，为保护旧结果不覆盖: {args.output_dir}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "legacy-whois-org-preprocess.json"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if not args.no_xlsx:
        builder = Path(__file__).with_name("render_legacy_org_preprocess_reports.mjs")
        subprocess.run(["node", str(builder), str(json_path), str(args.output_dir)], check=True)
    print(json.dumps({
        "status": report["status"], "summary": report["summary"],
        "name_stages": [stage["summary"] for stage in report["name_stages"]],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
