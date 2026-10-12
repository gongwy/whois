#!/usr/bin/env python3
"""Run the reproducible legacy-WHOIS + filing-data migration pipeline.

The current automated boundary is the organisation-resolution gate.  This
runner validates and discovers input workbooks, executes regression tests,
builds all canonical stage-1 outputs, creates the XLSX review workbook, and
verifies the saved workbook.  A blocked organisation stage still produces the
complete review package and exits with status 2 by design.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from openpyxl import load_workbook

from prepare_joint_rsdb_stage1 import (
    DOMAIN_REGISTRATION_REQUIRED_HEADERS,
    ICP_SHEET,
    IP_ALLOCATION_SHEET,
    REQUIRED_RPSL,
    Stage1,
    write_outputs,
)


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_DIR = SCRIPT_DIR.parent
REPORT_BUILDER = SCRIPT_DIR / "build_joint_stage1_report.mjs"
REPORT_FILENAME = "机构解析冲突报告.xlsx"
EXPECTED_REPORT_SHEETS = [
    "概览", "冲突明细", "机构候选", "匹配证据", "未解析域名",
    "未参与机构提取", "排除旧资源", "决策录入", "来源清单",
]


def workbook_sheets(path: Path) -> set[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"无法读取工作簿 {path}: {error}") from error
    sheets = set(workbook.sheetnames)
    workbook.close()
    return sheets


def discover_workbook(filing_dir: Path, sheet_name: str, explicit: Path | None) -> Path:
    if explicit is not None:
        path = explicit.resolve()
        if not path.is_file():
            raise FileNotFoundError(f"指定的备案工作簿不存在: {path}")
        if sheet_name not in workbook_sheets(path):
            raise ValueError(f"{path.name} 缺少工作表 {sheet_name}")
        return path

    matches = []
    for path in sorted(filing_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if sheet_name in workbook_sheets(path):
            matches.append(path.resolve())
    if not matches:
        raise FileNotFoundError(f"{filing_dir} 中没有包含工作表 {sheet_name} 的xlsx文件")
    if len(matches) > 1:
        names = "\n".join(f"- {path}" for path in matches)
        raise ValueError(
            f"发现多个包含工作表 {sheet_name} 的工作簿，请用显式参数指定:\n{names}"
        )
    return matches[0]


def workbook_has_headers(path: Path, required_headers: set[str]) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = {str(value).strip() for value in next(rows) if value is not None}
            except StopIteration:
                continue
            if required_headers.issubset(headers):
                return True
        return False
    finally:
        workbook.close()


def discover_domain_workbook(filing_dir: Path, explicit: Path | None) -> Path:
    if explicit is not None:
        path = explicit.resolve()
        if not path.is_file():
            raise FileNotFoundError(f"指定的域名备案工作簿不存在: {path}")
        if not workbook_has_headers(path, DOMAIN_REGISTRATION_REQUIRED_HEADERS):
            raise ValueError(f"{path.name} 缺少域名备案字段")
        return path
    matches = [
        path.resolve() for path in sorted(filing_dir.glob("*.xlsx"))
        if not path.name.startswith("~$")
        and workbook_has_headers(path, DOMAIN_REGISTRATION_REQUIRED_HEADERS)
    ]
    if len(matches) != 1:
        names = "\n".join(f"- {path}" for path in matches) or "- 无"
        raise ValueError(f"域名备案工作簿必须唯一，请用--domain-workbook指定:\n{names}")
    return matches[0]


def validate_rsdb(rsdb_dir: Path) -> None:
    if not rsdb_dir.is_dir():
        raise FileNotFoundError(f"旧WHOIS目录不存在: {rsdb_dir}")
    missing = [filename for filename, _ in REQUIRED_RPSL.values() if not (rsdb_dir / filename).is_file()]
    if missing:
        raise FileNotFoundError("旧WHOIS目录缺少文件: " + ", ".join(missing))


def run_regression_tests() -> None:
    command = [sys.executable, "-m", "unittest", "discover", "-v", "-s", "tools", "-p", "test_*.py"]
    completed = subprocess.run(command, cwd=REPO_DIR, check=False)
    if completed.returncode != 0:
        raise RuntimeError(f"机构阶段回归测试失败，退出码 {completed.returncode}")


def build_report(report_json: Path, output_xlsx: Path, preview_dir: Path) -> None:
    if not REPORT_BUILDER.is_file():
        raise FileNotFoundError(f"缺少报表生成脚本: {REPORT_BUILDER}")
    command = [
        "node", str(REPORT_BUILDER), str(report_json), str(output_xlsx), str(preview_dir),
    ]
    completed = subprocess.run(command, cwd=REPO_DIR, check=False)
    if completed.returncode != 0:
        raise RuntimeError(f"XLSX报告生成失败，退出码 {completed.returncode}")


def verify_outputs(output_dir: Path, report: dict, include_xlsx: bool) -> None:
    required = [
        "stage1-report.json", "source-manifest.json", "organisation-conflicts.tsv",
        "organisation-candidates.tsv", "organisation-evidence.tsv", "unresolved-domains.tsv",
        "non-organisation-filing-records.tsv", "excluded-filing-records.tsv",
        "excluded-legacy-resources.tsv",
        "organisation-decisions-template.tsv", "stage1-summary.md",
    ]
    if include_xlsx:
        required.append(REPORT_FILENAME)
    missing = [name for name in required if not (output_dir / name).is_file()]
    if missing:
        raise RuntimeError("输出不完整，缺少: " + ", ".join(missing))

    saved_report = json.loads((output_dir / "stage1-report.json").read_text(encoding="utf-8"))
    if saved_report["stats"] != report["stats"] or saved_report["status"] != report["status"]:
        raise RuntimeError("保存后的stage1报告与内存结果不一致")

    if not include_xlsx:
        return

    xlsx_path = output_dir / REPORT_FILENAME
    with zipfile.ZipFile(xlsx_path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise RuntimeError(f"XLSX压缩结构损坏: {bad_member}")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    if workbook.sheetnames != EXPECTED_REPORT_SHEETS:
        actual = ", ".join(workbook.sheetnames)
        workbook.close()
        raise RuntimeError(f"XLSX工作表结构异常: {actual}")
    workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rsdb-dir", type=Path, required=True, help="包含六个CERNIC.db.*旧WHOIS文件的目录")
    parser.add_argument("--filing-dir", type=Path, required=True, help="备案xlsx及可选排除规则所在目录")
    parser.add_argument("--icp-workbook", type=Path, help=f"可选；显式指定包含{ICP_SHEET}的工作簿")
    parser.add_argument("--domain-workbook", type=Path, help="可选；显式指定含name和registrant字段的域名备案工作簿")
    parser.add_argument("--ip-workbook", type=Path, help=f"可选；显式指定包含{IP_ALLOCATION_SHEET}的工作簿")
    parser.add_argument("--filing-exclusions", type=Path, help="可选；备案排除规则TSV")
    parser.add_argument("--output-dir", type=Path, required=True, help="必须为空或尚不存在")
    parser.add_argument("--skip-tests", action="store_true", help="跳过单元回归；正式迁移不建议使用")
    parser.add_argument("--skip-xlsx", action="store_true", help="只生成规范JSON/TSV，不生成审阅工作簿")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    rsdb_dir = args.rsdb_dir.resolve()
    filing_dir = args.filing_dir.resolve()
    output_dir = args.output_dir.resolve()
    if not filing_dir.is_dir():
        raise SystemExit(f"ERROR: 备案数据目录不存在: {filing_dir}")
    if output_dir.exists() and not output_dir.is_dir():
        raise SystemExit(f"ERROR: 输出路径不是目录: {output_dir}")
    if output_dir.exists() and any(output_dir.iterdir()):
        raise SystemExit(f"ERROR: 输出目录必须为空: {output_dir}")

    try:
        validate_rsdb(rsdb_dir)
        icp_workbook = discover_workbook(filing_dir, ICP_SHEET, args.icp_workbook)
        domain_workbook = discover_domain_workbook(filing_dir, args.domain_workbook)
        ip_workbook = discover_workbook(filing_dir, IP_ALLOCATION_SHEET, args.ip_workbook)
        if args.filing_exclusions:
            exclusions = args.filing_exclusions.resolve()
            if not exclusions.is_file():
                raise FileNotFoundError(f"指定的备案排除规则不存在: {exclusions}")
        else:
            default_exclusions = filing_dir / "迁移排除记录.tsv"
            exclusions = default_exclusions if default_exclusions.is_file() else None
        if not args.skip_tests:
            run_regression_tests()

        stage = Stage1(
            rsdb_dir, icp_workbook, ip_workbook, output_dir, exclusions,
            domain_workbook,
        )
        report = stage.run()
        write_outputs(output_dir, report)
        if not args.skip_xlsx:
            build_report(
                output_dir / "stage1-report.json",
                output_dir / REPORT_FILENAME,
                output_dir / "preview",
            )
        verify_outputs(output_dir, report, include_xlsx=not args.skip_xlsx)

        result = {
            "status": report["status"],
            "stage": report["stage"],
            "output_dir": str(output_dir),
            "rsdb_dir": str(rsdb_dir),
            "icp_workbook": str(icp_workbook),
            "domain_workbook": str(domain_workbook),
            "ip_workbook": str(ip_workbook),
            "filing_exclusions": str(exclusions) if exclusions else "",
            "conflicts": report["stats"].get("conflicts", 0),
            "blocking_conflicts": report["stats"].get("blocking_conflicts", 0),
            "candidates": report["stats"].get("organisation_candidates", 0),
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if report["status"] == "blocked":
            raise SystemExit(2)
    except SystemExit:
        raise
    except Exception as error:
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "pipeline-error.txt").write_text(
            f"{type(error).__name__}: {error}\n", encoding="utf-8",
        )
        raise SystemExit(f"ERROR: {error}") from error


if __name__ == "__main__":
    main()
