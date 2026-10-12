#!/usr/bin/env python3
"""Compare legal-entity data in ICP and IPFP filing workbooks.

The comparison deliberately ignores all IP and contact fields.  ICP contributes
domains and legal-entity fields; IPFP contributes legal-entity fields only.
IPFP rows whose allocation method is ``预留`` are retained for audit but do not
participate in matching or conflict detection.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
import json
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable

from openpyxl import load_workbook


ICP_SHEET = "ICP备案数据"
IPFP_SHEET = "IP分配信息"
PLACEHOLDERS = {"", "-", "--", "无", "暂无", "不详", "未知"}
DOMAIN_DOTS = str.maketrans({"。": ".", "．": ".", "｡": "."})
DOMAIN_SPLIT = re.compile(r"[\s,，;；、]+")
DOMAIN_LABEL = re.compile(r"[a-z0-9-]+")


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    normalized = unicodedata.normalize("NFKC", str(value)).strip()
    return "".join(
        character
        for character in normalized
        if unicodedata.category(character) not in {"Cf", "Cc"}
    ).strip()


def meaningful(value: object) -> str:
    candidate = text(value)
    return "" if candidate.casefold() in PLACEHOLDERS else candidate


def normalized(value: object) -> str:
    return re.sub(r"[\s\-_.()（）,，、/\\:：]+", "", text(value).casefold())


def normalized_address(value: object) -> str:
    return re.sub(r"[\s\-_.()（）,，、/\\:：号]+", "", text(value).casefold())


def province(value: object) -> str:
    first = re.split(r"[-—/]", meaningful(value), maxsplit=1)[0]
    return normalized(first)


def canonical_domain(value: object) -> str | None:
    candidate = text(value).translate(DOMAIN_DOTS).strip().rstrip(".")
    if not candidate or "_" in candidate:
        return None
    try:
        result = candidate.encode("idna").decode("ascii").lower()
    except UnicodeError:
        return None
    labels = result.split(".")
    if len(labels) < 2 or len(result) > 253:
        return None
    if any(
        not label
        or len(label) > 63
        or label.startswith("-")
        or label.endswith("-")
        or not DOMAIN_LABEL.fullmatch(label)
        for label in labels
    ):
        return None
    return result


def domain_tokens(value: object) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for token in DOMAIN_SPLIT.split(text(value).translate(DOMAIN_DOTS)):
        canonical = canonical_domain(token)
        if canonical and canonical not in seen:
            result.append(canonical)
            seen.add(canonical)
    return result


def read_sheet(path: Path, sheet_name: str) -> list[tuple[int, dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{path.name} 缺少工作表 {sheet_name}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(iterator)]
    except StopIteration:
        workbook.close()
        return []
    rows: list[tuple[int, dict[str, str]]] = []
    for excel_row, values in enumerate(iterator, 2):
        record = {
            header: text(values[index]) if index < len(values) else ""
            for index, header in enumerate(headers)
        }
        if any(record.values()):
            rows.append((excel_row, record))
    workbook.close()
    return rows


def require_headers(
    rows: list[tuple[int, dict[str, str]]], required: set[str], label: str
) -> None:
    headers = set(rows[0][1]) if rows else set()
    missing = sorted(required - headers)
    if missing:
        raise ValueError(f"{label} 缺少字段: {', '.join(missing)}")


def unique(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def identity_key(cert_type: str, cert_no: str) -> str:
    return f"{normalized(cert_type)}|{normalized(cert_no)}"


def identity_display(cert_type: str, cert_no: str) -> str:
    return f"{cert_type} / {cert_no}" if cert_type else cert_no


def compact_rows(rows: Iterable[int]) -> str:
    values = list(dict.fromkeys(rows))
    return "、".join(str(value) for value in values)


def joined(values: Iterable[str]) -> str:
    return "\n".join(unique(values))


def make_icp_records(rows: list[tuple[int, dict[str, str]]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for excel_row, row in rows:
        cert_type = meaningful(row["主办单位有效证件类型"])
        cert_no = meaningful(row["主办单位有效证件号码"])
        records.append(
            {
                "source_row": excel_row,
                "name": meaningful(row["主体单位名称"]),
                "nature": meaningful(row["主办单位性质"]),
                "cert_type": cert_type,
                "cert_no": cert_no,
                "identity_key": identity_key(cert_type, cert_no),
                "region": meaningful(row["主办单位所属区域"]),
                "province": province(row["主办单位所属区域"]),
                "registered_address": meaningful(row["主办单位证件住所"]),
                "mailing_address": meaningful(row["主办单位通讯地址"]),
                "domain_raw": meaningful(row["域名"]),
                "domains": domain_tokens(row["域名"]),
            }
        )
    return records


def make_ipfp_records(rows: list[tuple[int, dict[str, str]]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for excel_row, row in rows:
        cert_type = meaningful(row["使用单位证件类型"])
        cert_no = meaningful(row["使用单位证件号码"])
        allocation_method = meaningful(row["分配方式"])
        name = meaningful(row["使用单位名称"])
        records.append(
            {
                "source_row": excel_row,
                "record_id": meaningful(row["ID"]),
                "allocation_method": allocation_method,
                # A non-reserved row with a certificate remains comparable
                # even when its organisation name is missing.  This keeps the
                # missing-name issue visible without inventing a name.
                "included_in_comparison": allocation_method != "预留" and bool(cert_no),
                "name": name,
                "nature": meaningful(row["使用单位性质"]),
                "classification": meaningful(row["使用单位分类"]),
                "industry": meaningful(row["单位行业分类"]),
                "cert_type": cert_type,
                "cert_no": cert_no,
                "identity_key": identity_key(cert_type, cert_no),
                "region": meaningful(row["使用单位行政区域"]),
                "province": province(row["使用单位行政区域"]),
                "address": meaningful(row["使用单位详细地址"]),
            }
        )
    return records


def aggregate(records: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "names": unique(record["name"] for record in records),
        "natures": unique(record["nature"] for record in records),
        "cert_types": unique(record["cert_type"] for record in records),
        "cert_nos": unique(record["cert_no"] for record in records),
        "regions": unique(record["region"] for record in records),
        "provinces": unique(record["province"] for record in records),
        "addresses": unique(
            value
            for record in records
            for value in (
                record.get("registered_address", ""),
                record.get("mailing_address", ""),
                record.get("address", ""),
            )
        ),
        "domains": unique(
            domain for record in records for domain in record.get("domains", [])
        ),
        "source_rows": [record["source_row"] for record in records],
        "record_ids": unique(record.get("record_id", "") for record in records),
        "allocation_methods": unique(
            record.get("allocation_method", "") for record in records
        ),
        "record_count": len(records),
    }


def normalized_set(values: Iterable[str]) -> set[str]:
    return {normalized(value) for value in values if normalized(value)}


def address_set(values: Iterable[str]) -> set[str]:
    return {normalized_address(value) for value in values if normalized_address(value)}


def build_report(icp_path: Path, ipfp_path: Path) -> dict[str, Any]:
    icp_rows = read_sheet(icp_path, ICP_SHEET)
    ipfp_rows = read_sheet(ipfp_path, IPFP_SHEET)
    require_headers(
        icp_rows,
        {
            "主体单位名称",
            "主办单位性质",
            "主办单位有效证件类型",
            "主办单位有效证件号码",
            "主办单位所属区域",
            "主办单位证件住所",
            "主办单位通讯地址",
            "域名",
        },
        ICP_SHEET,
    )
    require_headers(
        ipfp_rows,
        {
            "ID",
            "分配方式",
            "使用单位名称",
            "使用单位性质",
            "使用单位分类",
            "单位行业分类",
            "使用单位行政区域",
            "使用单位详细地址",
            "使用单位证件类型",
            "使用单位证件号码",
        },
        IPFP_SHEET,
    )

    icp = make_icp_records(icp_rows)
    ipfp = make_ipfp_records(ipfp_rows)
    active_ipfp = [record for record in ipfp if record["included_in_comparison"]]
    reserved_named = [
        record
        for record in ipfp
        if record["allocation_method"] == "预留" and record["name"]
    ]

    icp_by_identity: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ipfp_by_identity: dict[str, list[dict[str, Any]]] = defaultdict(list)
    icp_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ipfp_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ipfp_by_cert_no: dict[str, list[dict[str, Any]]] = defaultdict(list)
    domain_map: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for record in icp:
        icp_by_identity[record["identity_key"]].append(record)
        icp_by_name[normalized(record["name"])].append(record)
        for domain in record["domains"]:
            domain_map[domain].append(record)
    for record in active_ipfp:
        ipfp_by_identity[record["identity_key"]].append(record)
        if normalized(record["name"]):
            ipfp_by_name[normalized(record["name"])].append(record)
        ipfp_by_cert_no[normalized(record["cert_no"])].append(record)

    conflicts: list[dict[str, Any]] = []
    conflict_counter = 0

    def add_conflict(
        severity: str,
        category: str,
        subtype: str,
        match_key: str,
        icp_records: list[dict[str, Any]],
        ipfp_records: list[dict[str, Any]],
        explanation: str,
        action: str,
    ) -> None:
        nonlocal conflict_counter
        conflict_counter += 1
        icp_values = aggregate(icp_records) if icp_records else aggregate([])
        ipfp_values = aggregate(ipfp_records) if ipfp_records else aggregate([])
        conflicts.append(
            {
                "conflict_id": f"FIL-C{conflict_counter:06d}",
                "severity": severity,
                "category": category,
                "subtype": subtype,
                "match_key": match_key,
                "icp_names": joined(icp_values["names"]),
                "ipfp_names": joined(ipfp_values["names"]),
                "icp_cert_types": joined(icp_values["cert_types"]),
                "ipfp_cert_types": joined(ipfp_values["cert_types"]),
                "icp_cert_nos": joined(icp_values["cert_nos"]),
                "ipfp_cert_nos": joined(ipfp_values["cert_nos"]),
                "icp_natures": joined(icp_values["natures"]),
                "ipfp_natures": joined(ipfp_values["natures"]),
                "icp_regions": joined(icp_values["regions"]),
                "ipfp_regions": joined(ipfp_values["regions"]),
                "icp_addresses": joined(icp_values["addresses"]),
                "ipfp_addresses": joined(ipfp_values["addresses"]),
                "icp_domains": joined(icp_values["domains"]),
                "icp_source_rows": compact_rows(icp_values["source_rows"]),
                "ipfp_source_rows": compact_rows(ipfp_values["source_rows"]),
                "ipfp_record_ids": joined(ipfp_values["record_ids"]),
                "explanation": explanation,
                "recommended_action": action,
            }
        )

    # ICP identity checks are included explicitly so future data refreshes do
    # not silently introduce internal ambiguities.
    for key in sorted(icp_by_identity):
        records = icp_by_identity[key]
        names = normalized_set(record["name"] for record in records)
        if len(names) > 1:
            add_conflict(
                "中",
                "ICP内部身份冲突",
                "ICP同一证件身份对应多个单位名称",
                key,
                records,
                ipfp_by_identity.get(key, []),
                "ICP 中相同的证件类型和证件号码对应多个不同单位名称。可能是更名，也可能是主体字段错误。",
                "核对证件号码对应的登记主体，并确认是否为历史名称。",
            )

    for name_key in sorted(icp_by_name):
        records = icp_by_name[name_key]
        identities = {record["identity_key"] for record in records}
        if len(identities) > 1:
            add_conflict(
                "中",
                "ICP内部身份冲突",
                "ICP同一单位名称对应多个证件身份",
                name_key,
                records,
                ipfp_by_name.get(name_key, []),
                "ICP 中同一规范化单位名称出现多个证件身份。可能是证件更新，也可能是证件号误填。",
                "核对证件有效期和登记主体，区分现行证件、历史证件与错误记录。",
            )

    # Same certificate identity should not point to several distinct legal names.
    for key in sorted(ipfp_by_identity):
        records = ipfp_by_identity[key]
        names = normalized_set(record["name"] for record in records)
        if len(names) > 1:
            add_conflict(
                "中",
                "IPFP内部身份冲突",
                "IPFP同一证件身份对应多个单位名称",
                key,
                icp_by_identity.get(key, []),
                records,
                "IPFP 中相同的证件类型和证件号码对应多个不同单位名称。可能是更名，也可能是行内容串接或证件误填。",
                "核对证件号码对应的登记主体，并标注历史名称；若名称互不相关，应修正错误来源行。",
            )

    # Same legal name should not carry several certificate identities without explanation.
    for name_key in sorted(ipfp_by_name):
        records = ipfp_by_name[name_key]
        identities = {record["identity_key"] for record in records}
        if len(identities) > 1:
            icp_records = icp_by_name.get(name_key, [])
            add_conflict(
                "中",
                "IPFP内部身份冲突",
                "IPFP同一单位名称对应多个证件身份",
                name_key,
                icp_records,
                records,
                "IPFP 中同一规范化单位名称出现多个证件类型或证件号码。可能是证件更新，也可能是证件号误填。",
                "核对每个证件的有效期和主体，保留现行证件并记录历史证件；错误值应回源修正。",
            )

    # A certificate number changing its type is an independent, high-value inconsistency.
    for cert_no_key in sorted(ipfp_by_cert_no):
        records = ipfp_by_cert_no[cert_no_key]
        cert_types = normalized_set(record["cert_type"] for record in records)
        if len(cert_types) > 1:
            related_icp = [
                record
                for key, values in icp_by_identity.items()
                if key.endswith("|" + cert_no_key)
                for record in values
            ]
            add_conflict(
                "高",
                "IPFP内部证件冲突",
                "IPFP同一证件号码对应多个证件类型",
                cert_no_key,
                related_icp,
                records,
                "IPFP 中相同证件号码被标成不同证件类型，不能直接确定唯一法人身份键。",
                "核对证件原件或权威登记信息，统一证件类型后再参与机构合并。",
            )

    # Exact names that disagree on identity are cross-source identity conflicts.
    for name_key in sorted(set(icp_by_name) & set(ipfp_by_name)):
        icp_records = icp_by_name[name_key]
        ipfp_records = ipfp_by_name[name_key]
        icp_identities = {record["identity_key"] for record in icp_records}
        ipfp_identities = {record["identity_key"] for record in ipfp_records}
        if icp_identities != ipfp_identities:
            add_conflict(
                "高",
                "跨源身份冲突",
                "跨源同名单位证件身份不一致",
                name_key,
                icp_records,
                ipfp_records,
                "ICP 与 IPFP 的单位名称一致，但证件类型或证件号码集合不同。",
                "以证件原件或权威登记为准确认当前证件；不要仅凭名称自动合并。",
            )

    # A domain assigned to multiple legal identities would be a direct ICP conflict.
    for domain in sorted(domain_map):
        records = domain_map[domain]
        identities = {record["identity_key"] for record in records}
        if len(identities) > 1:
            add_conflict(
                "高",
                "ICP内部域名冲突",
                "ICP同一完整域名对应多个法人身份",
                domain,
                records,
                [],
                "ICP 中同一完整规范化域名对应多个证件身份。域名按完整值比较，不折叠子域名。",
                "核对域名备案主体及备案状态，确认应保留的法人身份。",
            )

    common_identities = sorted(set(icp_by_identity) & set(ipfp_by_identity))
    differences: list[dict[str, Any]] = []
    matches: list[dict[str, Any]] = []
    difference_counter = 0

    def add_difference(
        identity: str,
        priority: str,
        subtype: str,
        field: str,
        icp_value: str,
        ipfp_value: str,
        icp_records: list[dict[str, Any]],
        ipfp_records: list[dict[str, Any]],
        explanation: str,
    ) -> None:
        nonlocal difference_counter
        difference_counter += 1
        differences.append(
            {
                "difference_id": f"FIL-D{difference_counter:06d}",
                "priority": priority,
                "subtype": subtype,
                "identity_key": identity,
                "field": field,
                "icp_value": icp_value,
                "ipfp_value": ipfp_value,
                "icp_source_rows": compact_rows(record["source_row"] for record in icp_records),
                "ipfp_source_rows": compact_rows(record["source_row"] for record in ipfp_records),
                "ipfp_record_ids": joined(record["record_id"] for record in ipfp_records),
                "explanation": explanation,
            }
        )

    for identity in common_identities:
        icp_records = icp_by_identity[identity]
        ipfp_records = ipfp_by_identity[identity]
        icp_values = aggregate(icp_records)
        ipfp_values = aggregate(ipfp_records)
        name_equal = normalized_set(icp_values["names"]) == normalized_set(ipfp_values["names"])
        nature_equal = normalized_set(icp_values["natures"]) == normalized_set(ipfp_values["natures"])
        province_equal = set(icp_values["provinces"]) == set(ipfp_values["provinces"])
        icp_address_keys = address_set(icp_values["addresses"])
        ipfp_address_keys = address_set(ipfp_values["addresses"])
        # The two sources use different address semantics (registered/mailing
        # versus resource-use address).  Treat an overlapping address as
        # consistent even when either source also carries additional campuses.
        address_equal = (
            (not icp_address_keys and not ipfp_address_keys)
            or bool(icp_address_keys & ipfp_address_keys)
        )
        if not name_equal:
            add_difference(
                identity,
                "中",
                "机构名称差异",
                "name",
                joined(icp_values["names"]),
                joined(ipfp_values["names"]),
                icp_records,
                ipfp_records,
                "证件身份相同但名称集合不同。常见原因包括更名、简称、旧名称或错误串接。",
            )
        if not nature_equal:
            add_difference(
                identity,
                "高",
                "单位性质差异",
                "org-type",
                joined(icp_values["natures"]),
                joined(ipfp_values["natures"]),
                icp_records,
                ipfp_records,
                "证件身份相同但单位性质不同，应核对登记类型。",
            )
        if not province_equal:
            add_difference(
                identity,
                "高",
                "省级行政区域差异",
                "province",
                joined(icp_values["regions"]),
                joined(ipfp_values["regions"]),
                icp_records,
                ipfp_records,
                "证件身份相同但来源记录出现不同省级行政区域，可能存在行内容串接或异地校区。",
            )
        if not address_equal:
            add_difference(
                identity,
                "低",
                "地址差异（非身份冲突）",
                "address",
                joined(icp_values["addresses"]),
                joined(ipfp_values["addresses"]),
                icp_records,
                ipfp_records,
                "ICP 的证件住所/通讯地址与 IPFP 的使用地址口径可能不同，仅列示，不据此拆分法人。",
            )
        matches.append(
            {
                "identity_key": identity,
                "cert_type": joined(icp_values["cert_types"] or ipfp_values["cert_types"]),
                "cert_no": joined(icp_values["cert_nos"] or ipfp_values["cert_nos"]),
                "icp_names": joined(icp_values["names"]),
                "ipfp_names": joined(ipfp_values["names"]),
                "icp_natures": joined(icp_values["natures"]),
                "ipfp_natures": joined(ipfp_values["natures"]),
                "icp_regions": joined(icp_values["regions"]),
                "ipfp_regions": joined(ipfp_values["regions"]),
                "icp_addresses": joined(icp_values["addresses"]),
                "ipfp_addresses": joined(ipfp_values["addresses"]),
                "icp_domains": joined(icp_values["domains"]),
                "name_status": "一致" if name_equal else "差异",
                "nature_status": "一致" if nature_equal else "差异",
                "province_status": "一致" if province_equal else "差异",
                "address_status": "一致" if address_equal else "差异",
                "icp_source_rows": compact_rows(icp_values["source_rows"]),
                "ipfp_source_rows": compact_rows(ipfp_values["source_rows"]),
                "ipfp_record_ids": joined(ipfp_values["record_ids"]),
            }
        )

    only_one_side: list[dict[str, Any]] = []
    for source, identities, own, other in (
        ("仅ICP", sorted(set(icp_by_identity) - set(ipfp_by_identity)), icp_by_identity, ipfp_by_identity),
        ("仅IPFP", sorted(set(ipfp_by_identity) - set(icp_by_identity)), ipfp_by_identity, icp_by_identity),
    ):
        del other
        for identity in identities:
            values = aggregate(own[identity])
            only_one_side.append(
                {
                    "source": source,
                    "identity_key": identity,
                    "names": joined(values["names"]),
                    "cert_types": joined(values["cert_types"]),
                    "cert_nos": joined(values["cert_nos"]),
                    "natures": joined(values["natures"]),
                    "regions": joined(values["regions"]),
                    "addresses": joined(values["addresses"]),
                    "domains": joined(values["domains"]),
                    "source_rows": compact_rows(values["source_rows"]),
                    "record_ids": joined(values["record_ids"]),
                    "note": "覆盖差异，不等同于数据冲突。",
                }
            )

    icp_domain_rows: list[dict[str, Any]] = []
    invalid_domain_rows = 0
    for record in icp:
        if not record["domains"]:
            invalid_domain_rows += 1
            tokens = [""]
        else:
            tokens = record["domains"]
        for domain in tokens:
            icp_domain_rows.append(
                {
                    "source_row": record["source_row"],
                    "name": record["name"],
                    "nature": record["nature"],
                    "cert_type": record["cert_type"],
                    "cert_no": record["cert_no"],
                    "region": record["region"],
                    "registered_address": record["registered_address"],
                    "mailing_address": record["mailing_address"],
                    "domain_raw": record["domain_raw"],
                    "domain_normalized": domain,
                }
            )

    ipfp_entity_rows = [
        {
            "source_row": record["source_row"],
            "record_id": record["record_id"],
            "allocation_method": record["allocation_method"],
            "included_in_comparison": "是" if record["included_in_comparison"] else "否",
            "name": record["name"],
            "nature": record["nature"],
            "classification": record["classification"],
            "industry": record["industry"],
            "cert_type": record["cert_type"],
            "cert_no": record["cert_no"],
            "region": record["region"],
            "address": record["address"],
        }
        for record in ipfp
        if record["name"] or record["included_in_comparison"]
    ]

    counts_by_subtype: dict[str, int] = defaultdict(int)
    for conflict in conflicts:
        counts_by_subtype[conflict["subtype"]] += 1
    differences_by_subtype: dict[str, int] = defaultdict(int)
    for difference in differences:
        differences_by_subtype[difference["subtype"]] += 1

    conflict_check_counts = {
        "ICP同一证件身份对应多个单位名称": sum(
            1
            for records in icp_by_identity.values()
            if len(normalized_set(record["name"] for record in records)) > 1
        ),
        "ICP同一单位名称对应多个证件身份": sum(
            1
            for records in icp_by_name.values()
            if len({record["identity_key"] for record in records}) > 1
        ),
        "ICP同一完整域名对应多个法人身份": sum(
            1
            for records in domain_map.values()
            if len({record["identity_key"] for record in records}) > 1
        ),
        "IPFP同一证件身份对应多个单位名称": sum(
            1
            for records in ipfp_by_identity.values()
            if len(normalized_set(record["name"] for record in records)) > 1
        ),
        "IPFP同一单位名称对应多个证件身份": sum(
            1
            for records in ipfp_by_name.values()
            if len({record["identity_key"] for record in records}) > 1
        ),
        "IPFP同一证件号码对应多个证件类型": sum(
            1
            for records in ipfp_by_cert_no.values()
            if len(normalized_set(record["cert_type"] for record in records)) > 1
        ),
        "跨源同名单位证件身份不一致": sum(
            1
            for name_key in set(icp_by_name) & set(ipfp_by_name)
            if {record["identity_key"] for record in icp_by_name[name_key]}
            != {record["identity_key"] for record in ipfp_by_name[name_key]}
        ),
    }

    return {
        "metadata": {
            "icp_source": str(icp_path),
            "ipfp_source": str(ipfp_path),
            "rules": [
                "ICP备案只提取完整域名和法人主体字段；忽略IP和所有联系人。",
                "IPFP只提取法人主体字段；IP字段不读取、不匹配、不判冲突。",
                "IPFP中分配方式为“预留”的记录保留在来源明细，但不参与本次匹配和冲突判定。",
                "主匹配键为规范化后的证件类型+证件号码；名称一致但证件身份不同不自动合并。",
                "域名按完整域名规范化，仅统一大小写、全角点和IDNA，不折叠www或子域名。",
                "地址差异单独列为非身份冲突；省级行政区域差异和单位性质差异需要优先复核。",
            ],
        },
        "summary": {
            "icp_source_rows": len(icp),
            "icp_domain_rows": len(icp_domain_rows),
            "icp_unique_domains": len(domain_map),
            "icp_unique_identities": len(icp_by_identity),
            "invalid_domain_rows": invalid_domain_rows,
            "ipfp_source_rows": len(ipfp),
            "ipfp_named_rows": sum(1 for record in ipfp if record["name"]),
            "ipfp_active_named_rows": sum(
                1 for record in active_ipfp if record["name"]
            ),
            "ipfp_comparison_rows": len(active_ipfp),
            "ipfp_reserved_named_rows": len(reserved_named),
            "ipfp_active_missing_name_rows": sum(
                1
                for record in ipfp
                if record["allocation_method"] != "预留" and not record["name"]
            ),
            "ipfp_unique_active_identities": len(ipfp_by_identity),
            "matched_identities": len(common_identities),
            "icp_only_identities": len(set(icp_by_identity) - set(ipfp_by_identity)),
            "ipfp_only_identities": len(set(ipfp_by_identity) - set(icp_by_identity)),
            "conflict_rows": len(conflicts),
            "difference_rows": len(differences),
            "conflict_check_counts": conflict_check_counts,
            "conflicts_by_subtype": dict(sorted(counts_by_subtype.items())),
            "differences_by_subtype": dict(sorted(differences_by_subtype.items())),
        },
        "conflicts": conflicts,
        "differences": differences,
        "matches": matches,
        "only_one_side": only_one_side,
        "icp_domain_rows": icp_domain_rows,
        "ipfp_entity_rows": ipfp_entity_rows,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--icp", type=Path, required=True)
    parser.add_argument("--ipfp", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = build_report(args.icp, args.ipfp)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
