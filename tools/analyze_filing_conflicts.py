#!/usr/bin/env python3
"""Build a complete, local-only conflict inventory for filing data vs WHOIS."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
import ipaddress
import json
from pathlib import Path
import re
import subprocess
import unicodedata

from openpyxl import load_workbook

import convert_rsdb_cernic as converter


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normal(value: object) -> str:
    value = unicodedata.normalize("NFKC", text(value)).lower()
    return re.sub(r"[\s\-_.()（）,，、/\\]+", "", value)


def normal_value(value: object) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text(value))).strip().casefold()


def latin_key(value: object) -> str:
    value = unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii").casefold()
    return re.sub(r"[^a-z0-9]", "", value)


def read_sheet(path: Path, sheet_name: str) -> list[tuple[int, dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows)]
    result = []
    for excel_row, values in enumerate(rows, 2):
        row = {
            headers[index]: text(values[index]) if index < len(values) else ""
            for index in range(len(headers))
        }
        if any(row.values()):
            result.append((excel_row, row))
    workbook.close()
    return result


def parse_dump(path: Path) -> list[dict[str, list[str]]]:
    objects = []
    current = []
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if not raw.strip():
            if current:
                objects.append(dict(converter.parse(current)))
                current = []
        else:
            current.append(raw)
    if current:
        objects.append(dict(converter.parse(current)))
    return objects


def object_type(obj: dict[str, list[str]]) -> str | None:
    for name in ("organisation", "person", "role", "inetnum", "inet6num", "domain"):
        if name in obj:
            return name
    return None


def one(obj: dict[str, list[str]], key: str) -> str:
    return obj.get(key, [""])[0]


DOMAIN_DOTS = str.maketrans({"。": ".", "．": ".", "｡": "."})
DOMAIN_CANDIDATE = re.compile(r"(?<![\w.-])(?:[\w-]+\.)+[\w-]+\.?(?![\w.-])", re.UNICODE)
ASCII_DOMAIN_LABEL = re.compile(r"[a-z0-9-]+")


def canonical_domain(value: str) -> str | None:
    """Return a lowercase ASCII/IDNA domain, or None when the value is invalid."""
    candidate = unicodedata.normalize("NFKC", value).translate(DOMAIN_DOTS).strip().rstrip(".")
    if not candidate or "_" in candidate:
        return None
    try:
        canonical = candidate.encode("idna").decode("ascii").lower()
    except UnicodeError:
        return None
    labels = canonical.split(".")
    if len(labels) < 2 or len(canonical) > 253:
        return None
    if any(
        not label
        or len(label) > 63
        or label.startswith("-")
        or label.endswith("-")
        or not ASCII_DOMAIN_LABEL.fullmatch(label)
        for label in labels
    ):
        return None
    tld = labels[-1]
    if not (tld.startswith("xn--") or re.fullmatch(r"[a-z]{2,63}", tld)):
        return None
    return canonical


def domain_tokens(value: str) -> list[tuple[str, str]]:
    """Extract domains as (source spelling, lowercase ASCII/IDNA spelling)."""
    normalized = unicodedata.normalize("NFKC", text(value)).translate(DOMAIN_DOTS)
    result: list[tuple[str, str]] = []
    seen: set[str] = set()
    for match in DOMAIN_CANDIDATE.finditer(normalized):
        original = match.group(0).rstrip(".")
        canonical = canonical_domain(original)
        if canonical and canonical not in seen:
            result.append((original, canonical))
            seen.add(canonical)
    return result


@dataclass(frozen=True)
class IpRange:
    version: int
    first: int
    last: int


def ip_range(start: str, end: str = "") -> IpRange | None:
    try:
        if not end and "/" in start:
            network = ipaddress.ip_network(start, strict=False)
            return IpRange(network.version, int(network.network_address), int(network.broadcast_address))
        first_ip = ipaddress.ip_address(start.strip())
        last_ip = ipaddress.ip_address(end.strip())
        if first_ip.version != last_ip.version or int(first_ip) > int(last_ip):
            return None
        return IpRange(first_ip.version, int(first_ip), int(last_ip))
    except ValueError:
        return None


def object_range(obj: dict[str, list[str]], typ: str) -> IpRange | None:
    value = one(obj, typ)
    if typ == "inet6num" and "/" in value:
        return ip_range(value)
    pieces = re.split(r"\s*-\s*", value, maxsplit=1)
    return ip_range(*pieces) if len(pieces) == 2 else None


def normal_phone(value: str) -> str:
    raw = text(value)
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return ""
    if raw.startswith("+"):
        candidate = "+" + digits
    elif len(digits) == 11 and digits.startswith("1"):
        candidate = "+86" + digits
    elif digits.startswith("86") and len(digits) >= 10:
        candidate = "+" + digits
    elif 7 <= len(digits) <= 18:
        candidate = "+86" + digits
    else:
        return ""
    return candidate if len(candidate) <= 30 else ""


EMAIL = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


@dataclass(frozen=True)
class Contact:
    name: str
    phones: tuple[str, ...]
    emails: tuple[str, ...]
    org: str

    @property
    def identity(self) -> tuple[str, tuple[str, ...], tuple[str, ...], str]:
        return normal_value(self.name), self.phones, self.emails, self.org

    def display(self) -> str:
        parts = [self.name]
        if self.phones:
            parts.append("电话=" + ", ".join(self.phones))
        if self.emails:
            parts.append("邮箱=" + ", ".join(self.emails))
        return "; ".join(parts)


def contact(name: str, phones: list[str], emails: list[str], org: str) -> Contact | None:
    name = text(name)
    if not name:
        return None
    valid_phones = tuple(dict.fromkeys(phone for phone in (normal_phone(value) for value in phones) if phone))
    valid_emails = tuple(dict.fromkeys(text(value).lower() for value in emails if EMAIL.fullmatch(text(value))))
    return Contact(name, valid_phones, valid_emails, org)


def transliterate(names: list[str]) -> dict[str, str]:
    names = list(dict.fromkeys(name for name in names if name))
    if not names:
        return {}
    script = '''ObjC.import("Foundation");
const inputData = $.NSFileHandle.fileHandleWithStandardInput.readDataToEndOfFile;
const input = ObjC.unwrap($.NSString.alloc.initWithDataEncoding(inputData, $.NSUTF8StringEncoding));
const values = JSON.parse(input);
const output = values.map((value) => {
  const mutable = $.NSMutableString.alloc.initWithString(value);
  $.CFStringTransform(mutable, null, $.kCFStringTransformToLatin, false);
  $.CFStringTransform(mutable, null, $.kCFStringTransformStripDiacritics, false);
  return ObjC.unwrap(mutable);
});
console.log(JSON.stringify(output));'''
    try:
        completed = subprocess.run(
            ["osascript", "-l", "JavaScript", "-e", script],
            input=json.dumps(names, ensure_ascii=False), text=True,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True, timeout=30,
        )
        values = json.loads(completed.stderr.strip().splitlines()[-1])
        if len(values) != len(names):
            return {}
        return {normal(name): latin_key(value) for name, value in zip(names, values)}
    except (OSError, subprocess.SubprocessError, ValueError, json.JSONDecodeError, IndexError):
        return {}


class Analyzer:
    def __init__(self, objects: list[dict[str, list[str]]]) -> None:
        self.objects = objects
        self.by_target: dict[tuple[str, str], dict[str, list[str]]] = {}
        self.org_names: dict[str, set[str]] = defaultdict(set)
        self.domains: dict[str, tuple[str, str]] = {}
        self.ranges: dict[int, list[tuple[IpRange, tuple[str, str]]]] = defaultdict(list)
        self.exact_ranges: dict[IpRange, list[tuple[str, str]]] = defaultdict(list)
        self.contacts: dict[str, dict[str, list[str]]] = {}
        self.pending_values: dict[tuple[tuple[str, str], str], list[dict]] = defaultdict(list)
        self.pending_contacts: dict[tuple[tuple[str, str], str], list[dict]] = defaultdict(list)
        self.issues: list[dict] = []
        self.stats = Counter()
        self._index()

    def _index(self) -> None:
        for obj in self.objects:
            typ = object_type(obj)
            if typ in {"person", "role"} and one(obj, "nic-hdl"):
                self.contacts[one(obj, "nic-hdl").upper()] = obj
                continue
            if typ not in {"organisation", "domain", "inetnum", "inet6num"}:
                continue
            key = one(obj, typ)
            target = (typ, key)
            self.by_target[target] = obj
            if typ == "organisation":
                aliases = list(obj.get("org-name", [])) + list(obj.get("c-org-name", []))
                aliases.extend(
                    value.split(":", 1)[1].strip()
                    for value in obj.get("descr", [])
                    if value.startswith("legacy-org-name:")
                )
                aliases.extend(
                    value.split(":", 1)[1].strip()
                    for value in obj.get("remarks", [])
                    if value.startswith("legacy-c-org-name:")
                )
                for alias in aliases:
                    if normal(alias):
                        self.org_names[normal(alias)].add(key)
            elif typ == "domain":
                self.domains[key.lower().rstrip(".")] = target
            else:
                interval = object_range(obj, typ)
                if interval:
                    self.ranges[interval.version].append((interval, target))
                    self.exact_ranges[interval].append(target)

    def add_issue(
        self, major: str, category: str, source: str, source_rows: str,
        target_type: str = "", target_key: str = "", match_basis: str = "",
        field: str = "", whois_value: str = "", filing_value: str = "",
        record_key: str = "", explanation: str = "", action: str = "",
    ) -> None:
        self.issues.append({
            "major_type": major,
            "conflict_type": category,
            "source_dataset": source,
            "source_rows": source_rows,
            "target_type": target_type,
            "target_key": target_key,
            "match_basis": match_basis,
            "field": field,
            "whois_value": whois_value,
            "filing_value": filing_value,
            "record_key": record_key,
            "explanation": explanation,
            "suggested_action": action,
        })

    def org_target(self, name: str, source: str, row: int, record_key: str) -> tuple[str, str] | None:
        if not text(name):
            self.add_issue(
                "匹配问题", "机构名称缺失", source, str(row),
                match_basis="备案单位名称", record_key=record_key,
                explanation="备案记录没有主体单位名称或使用单位名称。",
                action="补齐单位名称后重新匹配。",
            )
            self.stats["organisation_name_missing"] += 1
            return None
        matched = sorted(self.org_names.get(normal(name), set()))
        if len(matched) == 1:
            self.stats["organisation_unique_match"] += 1
            return "organisation", matched[0]
        if matched:
            self.add_issue(
                "匹配问题", "机构名称匹配多个WHOIS机构", source, str(row),
                "organisation", "\n".join(matched), "org-name/c-org-name/历史名称",
                filing_value=name, record_key=record_key,
                explanation="同一备案单位名称对应多个WHOIS机构，不能自动选择。",
                action="核对机构后指定唯一organisation。",
            )
            self.stats["organisation_ambiguous"] += 1
        else:
            self.add_issue(
                "匹配问题", "机构名称未匹配", source, str(row),
                match_basis="org-name/c-org-name/历史名称", filing_value=name,
                record_key=record_key,
                explanation="备案单位名称在WHOIS机构名称及历史名称中没有匹配。",
                action="确认是否新建机构、关联FK临时机构或忽略。",
            )
            self.stats["organisation_unmatched"] += 1
        return None

    def domain_target(
        self, domain: str, source: str, row: int, record_key: str,
        original_domain: str = "",
    ) -> tuple[str, str] | None:
        target = self.domains.get(domain.lower().rstrip("."))
        if target:
            self.stats["domain_exact_match"] += 1
            if original_domain and original_domain.casefold() != domain.casefold():
                self.stats["domain_idn_match"] += 1
            return target
        filing_value = original_domain or domain
        match_basis = "domain精确匹配"
        explanation = "备案域名在当前WHOIS domain对象中不存在。"
        action = "确认域名是否应新建、已注销或格式需要修正。"
        if original_domain and original_domain.casefold() != domain.casefold():
            filing_value = f"{original_domain}\nPunycode: {domain}"
            match_basis = "domain精确匹配（IDNA规范化）"
            explanation = "备案国际化域名已成功转换为Punycode，但当前WHOIS中不存在对应domain对象。"
            action = "确认是否应新建对应Punycode domain对象，或该域名已注销、不纳入WHOIS范围。"
            self.stats["domain_idn_unmatched"] += 1
        self.add_issue(
            "匹配问题", "域名未匹配", source, str(row), "domain", domain.upper(), match_basis,
            filing_value=filing_value, record_key=record_key,
            explanation=explanation,
            action=action,
        )
        self.stats["domain_unmatched"] += 1
        return None

    def resource_target(
        self, start: str, end: str, source: str, row: int, record_key: str,
    ) -> tuple[str, str] | None:
        if not start or not end:
            self.add_issue(
                "匹配问题", "IP区间缺失", source, str(row), match_basis="IP起止地址",
                filing_value=f"{start} - {end}", record_key=record_key,
                explanation="备案记录缺少起始IP或终止IP。",
                action="补齐IP区间后重新匹配。",
            )
            self.stats["ip_missing"] += 1
            return None
        interval = ip_range(start, end)
        if interval is None:
            self.add_issue(
                "匹配问题", "IP区间格式无效", source, str(row), match_basis="IP起止地址",
                filing_value=f"{start} - {end}", record_key=record_key,
                explanation="起止IP无法解析、版本不一致或起始地址大于终止地址。",
                action="修正备案IP区间。",
            )
            self.stats["ip_invalid"] += 1
            return None
        exact = self.exact_ranges.get(interval, [])
        if len(exact) == 1:
            self.stats["ip_exact_match"] += 1
            return exact[0]
        if len(exact) > 1:
            matches = exact
            basis = "IP区间完全匹配"
        else:
            matches = [
                target for current, target in self.ranges[interval.version]
                if current.first <= interval.first and interval.last <= current.last
            ]
            basis = "IP区间包含匹配"
        if len(matches) == 1:
            self.stats["ip_containment_match"] += 1
            return matches[0]
        category = "IP区间匹配多个WHOIS网段" if matches else "IP区间未匹配"
        self.add_issue(
            "匹配问题", category, source, str(row),
            target_key="\n".join(target[1] for target in matches), match_basis=basis,
            filing_value=f"{start} - {end}", record_key=record_key,
            explanation=("备案IP区间同时落入多个WHOIS网段。" if matches else "备案IP区间未找到完全匹配或唯一包含的WHOIS网段。"),
            action=("核对后指定唯一网段。" if matches else "确认网段是否缺失、已删除或边界需要修正。"),
        )
        self.stats["ip_ambiguous" if matches else "ip_unmatched"] += 1
        return None

    def add_value(self, target: tuple[str, str] | None, field: str, value: str, source: str, row: int) -> None:
        if target and text(value):
            self.pending_values[(target, field)].append({"source": source, "row": row, "value": text(value)})

    def add_contact(self, target: tuple[str, str] | None, field: str, value: Contact | None, source: str, row: int) -> None:
        if target and value:
            self.pending_contacts[(target, field)].append({"source": source, "row": row, "contact": value})

    def existing_contact_display(self, handles: list[str]) -> str:
        rows = []
        for handle in handles:
            obj = self.contacts.get(handle.upper())
            names = [] if obj is None else obj.get("c-name", []) + obj.get("person", [])
            rows.append(f"{handle}: {', '.join(dict.fromkeys(names))}" if names else handle)
        return "\n".join(rows)

    def contact_equivalent(self, incoming: Contact, handle: str, pinyin: dict[str, str]) -> bool:
        existing = self.contacts.get(handle.upper())
        if not existing:
            return False
        incoming_key = normal(incoming.name)
        existing_names = existing.get("c-name", []) + existing.get("person", [])
        if any(incoming_key and incoming_key == normal(name) for name in existing_names):
            return True
        translated = pinyin.get(incoming_key, "")
        return bool(translated and any(latin_key(name) == translated for name in existing_names))

    @staticmethod
    def source_label(entries: list[dict]) -> str:
        return "+".join(sorted({entry["source"] for entry in entries}))

    @staticmethod
    def source_rows(entries: list[dict]) -> str:
        return "\n".join(f"{entry['source']}!{entry['row']}" for entry in entries)

    @staticmethod
    def filing_values(entries: list[dict], key: str) -> str:
        values = []
        for entry in entries:
            value = entry[key]
            display = value.display() if isinstance(value, Contact) else value
            line = f"{entry['source']}!{entry['row']}: {display}"
            if line not in values:
                values.append(line)
        return "\n".join(values)

    def resolve(self) -> None:
        for (target, field), entries in self.pending_values.items():
            obj = self.by_target[target]
            original = obj.get(field, [])
            unique = {}
            for entry in entries:
                unique.setdefault(normal_value(entry["value"]), entry)
            if original:
                original_keys = {normal_value(value) for value in original if normal_value(value)}
                if not set(unique).issubset(original_keys):
                    self.add_issue(
                        "值冲突", "WHOIS与备案字段不一致", self.source_label(entries), self.source_rows(entries),
                        target[0], target[1], "已匹配对象", field, "\n".join(original),
                        self.filing_values(entries, "value"), target[1],
                        "WHOIS已有值，备案数据包含不同值；转换规则不会自动覆盖。",
                        "核对后决定保留WHOIS值或更新为备案值。",
                    )
            elif field != "c-address" and len(unique) > 1:
                self.add_issue(
                    "值冲突", "备案内部字段多值", self.source_label(entries), self.source_rows(entries),
                    target[0], target[1], "已匹配对象", field, "",
                    self.filing_values(entries, "value"), target[1],
                    "同一WHOIS对象的空字段收到多个不同备案候选值。",
                    "确认唯一值后再补充。",
                )

        incoming_names = [
            entry["contact"].name
            for entries in self.pending_contacts.values()
            for entry in entries
        ]
        pinyin = transliterate(incoming_names)
        self.stats["contact_names_transliterated"] = len(pinyin)
        for (target, field), entries in self.pending_contacts.items():
            obj = self.by_target[target]
            existing = obj.get(field, [])
            identities = {}
            for entry in entries:
                identities.setdefault(entry["contact"].identity, entry)
            unique_entries = list(identities.values())
            if existing:
                unmatched = [
                    entry for entry in unique_entries
                    if not any(self.contact_equivalent(entry["contact"], handle, pinyin) for handle in existing)
                ]
                if unmatched:
                    self.add_issue(
                        "值冲突", "WHOIS与备案联系人不一致", self.source_label(unmatched), self.source_rows(unmatched),
                        target[0], target[1], "已匹配对象", field,
                        self.existing_contact_display(existing), self.filing_values(unmatched, "contact"), target[1],
                        "WHOIS已有联系人，备案负责人姓名与现有person/role不等价。",
                        "核对负责人；确认后保留、替换或新增联系人。",
                    )
                else:
                    self.stats["contact_existing_equivalent"] += 1
            elif len(unique_entries) > 1:
                self.add_issue(
                    "值冲突", "备案内部联系人多值", self.source_label(unique_entries), self.source_rows(unique_entries),
                    target[0], target[1], "已匹配对象", field, "",
                    self.filing_values(unique_entries, "contact"), target[1],
                    "同一WHOIS对象的空联系人字段收到多个不同备案联系人。",
                    "确认唯一联系人后再新建person并关联。",
                )


def target_org(analyzer: Analyzer, target: tuple[str, str] | None) -> str:
    if not target:
        return ""
    return one(analyzer.by_target[target], "org")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--whois-dump", type=Path, required=True)
    parser.add_argument("--icp-workbook", type=Path, required=True)
    parser.add_argument("--ip-workbook", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    args = parser.parse_args()

    objects = parse_dump(args.whois_dump)
    analyzer = Analyzer(objects)
    icp_rows = read_sheet(args.icp_workbook, "ICP备案数据")
    ip_rows = read_sheet(args.ip_workbook, "IP分配信息")

    for excel_row, row in icp_rows:
        source = "ICP备案数据"
        record_key = row["主体单位名称"] or f"ICP备案数据第{excel_row}行"
        org = analyzer.org_target(row["主体单位名称"], source, excel_row, record_key)
        analyzer.add_value(org, "org-type", row["主办单位性质"], source, excel_row)
        analyzer.add_value(org, "cert-type", row["主办单位有效证件类型"], source, excel_row)
        analyzer.add_value(org, "cert-no", row["主办单位有效证件号码"], source, excel_row)
        for address in (row["主办单位证件住所"], row["主办单位通讯地址"]):
            analyzer.add_value(org, "c-address", address, source, excel_row)
        subject = contact(
            row["主体负责人姓名"],
            [row["主体负责人手机号码"], row["主体负责人办公电话"], row["主体负责人应急联系电话"]],
            [row["主体负责人邮件地址"]], org[1] if org else "",
        )
        service = contact(
            row["互联网服务负责人姓名"],
            [row["互联网服务负责人手机号码"], row["互联网服务负责人办公电话"], row["互联网服务负责人应急联系电话"]],
            [row["互联网服务负责人邮件地址"]], org[1] if org else "",
        )
        analyzer.add_contact(org, "admin-c", subject, source, excel_row)
        raw_domain = row["域名"]
        tokens = domain_tokens(raw_domain)
        if raw_domain and not tokens:
            analyzer.add_issue(
                "匹配问题", "域名格式无法解析", source, str(excel_row), filing_value=raw_domain,
                record_key=record_key, match_basis="域名字段",
                explanation="域名字段非空，但未能提取合法域名。",
                action="修正域名分隔符或格式。",
            )
        for original_domain, domain in tokens:
            target = analyzer.domain_target(domain, source, excel_row, record_key, original_domain)
            analyzer.add_contact(target, "admin-c", subject, source, excel_row)
            analyzer.add_contact(target, "tech-c", service, source, excel_row)
        resource = analyzer.resource_target(
            row["起始IP地址"], row["终止IP地址"], source, excel_row, record_key)
        resource_subject = contact(
            row["主体负责人姓名"],
            [row["主体负责人手机号码"], row["主体负责人办公电话"], row["主体负责人应急联系电话"]],
            [row["主体负责人邮件地址"]], target_org(analyzer, resource),
        )
        analyzer.add_contact(resource, "admin-c", resource_subject, source, excel_row)

    for excel_row, row in ip_rows:
        source = "IP分配信息"
        record_key = row["使用单位名称"] or f"IP分配信息第{excel_row}行"
        org = analyzer.org_target(row["使用单位名称"], source, excel_row, record_key)
        analyzer.add_value(org, "org-type", row["使用单位性质"], source, excel_row)
        analyzer.add_value(org, "c-address", row["使用单位详细地址"], source, excel_row)
        resource = analyzer.resource_target(row["起始IP"], row["终止IP"], source, excel_row, record_key)
        resource_contact = contact(
            row["联系人姓名"], [row["联系人电话"]], [row["联系人电子邮件"]],
            target_org(analyzer, resource),
        )
        analyzer.add_contact(resource, "tech-c", resource_contact, source, excel_row)

    analyzer.resolve()
    analyzer.issues.sort(key=lambda item: (
        item["major_type"], item["conflict_type"], item["target_type"],
        item["target_key"], item["source_dataset"], item["source_rows"],
    ))
    for index, issue in enumerate(analyzer.issues, 1):
        issue["conflict_id"] = f"C{index:06d}"

    category_counts = Counter(issue["conflict_type"] for issue in analyzer.issues)
    major_counts = Counter(issue["major_type"] for issue in analyzer.issues)
    representatives = []
    for category in sorted(category_counts):
        candidates = [issue for issue in analyzer.issues if issue["conflict_type"] == category]
        candidates.sort(key=lambda item: len(item["whois_value"]) + len(item["filing_value"]))
        representatives.extend(candidates[:3])

    result = {
        "source": {
            "whois_dump": args.whois_dump.name,
            "whois_objects": len(objects),
            "icp_workbook": args.icp_workbook.name,
            "icp_rows": len(icp_rows),
            "ip_workbook": args.ip_workbook.name,
            "ip_rows": len(ip_rows),
        },
        "major_counts": dict(sorted(major_counts.items())),
        "category_counts": dict(sorted(category_counts.items())),
        "match_stats": dict(sorted(analyzer.stats.items())),
        "issues": analyzer.issues,
        "representatives": representatives,
    }
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({
        "issues": len(analyzer.issues),
        "categories": len(category_counts),
        "major_counts": result["major_counts"],
        "category_counts": result["category_counts"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
