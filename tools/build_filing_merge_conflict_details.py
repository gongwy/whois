#!/usr/bin/env python3
"""Reconstruct row-level evidence for the legacy filing merge conflict report.

The legacy report aggregates candidate values per WHOIS organisation.  This
script replays the original name matching against the same baseline and emits
JSON with one row per conflict field plus every contributing filing source row.
It is intentionally read-only with respect to all source workbooks and RPSL.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
import importlib.util
import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook


def load_legacy_module(path: Path):
    spec = importlib.util.spec_from_file_location("legacy_filing_merger", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load legacy merger: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(path: Path, sheet_name: str) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(iterator)]
    rows: list[dict[str, str]] = []
    for excel_row, values in enumerate(iterator, start=2):
        record = {
            headers[index]: clean(values[index]) if index < len(values) else ""
            for index in range(len(headers))
        }
        if any(record.values()):
            record["__row__"] = str(excel_row)
            rows.append(record)
    workbook.close()
    return rows


def read_legacy_conflicts(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook["合并矛盾"]
    result: list[dict[str, str]] = []
    for row in range(4, sheet.max_row + 1):
        key = clean(sheet.cell(row, 1).value)
        typ = clean(sheet.cell(row, 2).value)
        fields = clean(sheet.cell(row, 3).value).splitlines()
        if not key or not typ:
            continue
        for field in fields:
            if field:
                result.append({
                    "object_key": key,
                    "object_type": typ,
                    "field": field,
                    "legacy_row": str(row),
                })
    workbook.close()
    return result


def read_legacy_report_rows(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook["合并矛盾"]
    rows = [
        [clean(sheet.cell(row, column).value) for column in range(1, 6)]
        for row in range(1, sheet.max_row + 1)
    ]
    workbook.close()
    return rows


def contact_text(obj: dict[str, list[str]] | None) -> str:
    if not obj:
        return ""
    parts: list[str] = []
    for field, label in (("person", "person"), ("c-name", "c-name"), ("phone", "phone"), ("e-mail", "e-mail")):
        values = obj.get(field, [])
        if values:
            parts.append(f"{label}: {', '.join(values)}")
    return "\n".join(parts)


def raw_object(obj: dict[str, list[str]]) -> str:
    return "\n".join(f"{field}: {value}" for field, values in obj.items() for value in values)


def source_display(source: dict[str, str], field: str) -> str:
    if field == "admin-c":
        details = [source.get("candidate", "")]
        if source.get("phones"):
            details.append(f"电话: {source['phones']}")
        if source.get("emails"):
            details.append(f"邮箱: {source['emails']}")
        return "；".join(part for part in details if part)
    return source.get("candidate", "")


def explanation(field: str, original: str, candidates: list[str]) -> tuple[str, str]:
    count = len(candidates)
    if field == "org-type":
        reason = f"同一 WHOIS 机构收到 {count} 个不同的单位性质候选值。"
        action = "核对来源行中的单位名称和单位性质；若实际为不同主体，应拆分或修正机构关联。"
    elif field == "cert-no":
        reason = f"同一 WHOIS 机构收到 {count} 个不同证件号码，可能把不同法人主体关联到同一机构。"
        action = "优先核验证件号码、主体单位名称和地址；确认是否需要拆分机构或修正备案记录。"
    elif field == "admin-c":
        reason = f"同一 WHOIS 机构收到 {count} 个不同主体负责人，无法唯一确定 admin-c。"
        action = "核对负责人姓名、电话、邮箱及对应主体；确认唯一联系人，或先拆分错误合并的机构。"
    else:
        reason = f"备案候选值与 WHOIS 原值不一致，或存在 {count} 个不同候选值。"
        action = "逐行核对来源记录，确认应保留的值及机构关联。"
    if original:
        reason += " WHOIS 当前已有值，因此旧合并流程未覆盖。"
    else:
        reason += " WHOIS 当前为空，但候选值不唯一，因此旧合并流程未自动写入。"
    return reason, action


def build(args: argparse.Namespace) -> dict[str, Any]:
    legacy = load_legacy_module(args.legacy_script)
    objects = legacy.parse_rpsl(args.baseline)
    merger = legacy.Merger(objects)
    legacy_conflicts = read_legacy_conflicts(args.original_report)
    conflict_keys = {(item["object_type"], item["object_key"], item["field"]) for item in legacy_conflicts}

    contributions: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)

    for row in read_rows(args.icp_workbook, "ICP备案数据"):
        target = merger.org_target(row.get("主体单位名称", ""))
        if not target:
            continue
        base = {
            "source_sheet": "ICP备案数据",
            "source_row": row["__row__"],
            "unit_name": row.get("主体单位名称", ""),
            "unit_type": row.get("主办单位性质", ""),
            "cert_type": row.get("主办单位有效证件类型", ""),
            "cert_no": row.get("主办单位有效证件号码", ""),
            "region": row.get("主办单位所属区域", ""),
            "address": "\n".join(dict.fromkeys(value for value in (
                row.get("主办单位证件住所", ""), row.get("主办单位通讯地址", "")
            ) if value)),
            "subject_contact": row.get("主体负责人姓名", ""),
            "phones": "、".join(dict.fromkeys(value for value in (
                row.get("主体负责人手机号码", ""), row.get("主体负责人办公电话", ""),
                row.get("主体负责人应急联系电话", "")
            ) if value)),
            "emails": row.get("主体负责人邮件地址", ""),
            "domain": row.get("域名", ""),
            "start_ip": row.get("起始IP地址", ""),
            "end_ip": row.get("终止IP地址", ""),
            "source_record_id": row.get("接入ID", ""),
        }
        for field, candidate in (
            ("org-type", row.get("主办单位性质", "")),
            ("cert-no", row.get("主办单位有效证件号码", "")),
            ("admin-c", row.get("主体负责人姓名", "")),
        ):
            key = (target[0], target[1], field)
            if key in conflict_keys and clean(candidate):
                contributions[key].append({**base, "candidate": clean(candidate)})

    for row in read_rows(args.ip_workbook, "IP分配信息"):
        target = merger.org_target(row.get("使用单位名称", ""))
        if not target:
            continue
        key = (target[0], target[1], "org-type")
        candidate = row.get("使用单位性质", "")
        if key in conflict_keys and clean(candidate):
            contributions[key].append({
                "source_sheet": "IP分配信息",
                "source_row": row["__row__"],
                "source_record_id": row.get("ID", ""),
                "unit_name": row.get("使用单位名称", ""),
                "unit_type": row.get("使用单位性质", ""),
                "candidate": clean(candidate),
                "cert_type": row.get("使用单位证件类型", ""),
                "cert_no": row.get("使用单位证件号码", ""),
                "region": row.get("使用单位行政区域", ""),
                "address": row.get("使用单位详细地址", ""),
                "subject_contact": row.get("联系人姓名", ""),
                "phones": row.get("联系人电话", ""),
                "emails": row.get("联系人电子邮件", ""),
                "domain": "",
                "start_ip": row.get("起始IP", ""),
                "end_ip": row.get("终止IP", ""),
            })

    contact_objects: dict[str, dict[str, list[str]]] = {}
    for obj in objects:
        if legacy.object_type(obj) in {"person", "role"} and obj.get("nic-hdl"):
            contact_objects[obj["nic-hdl"][0]] = obj

    resources_by_org: dict[str, list[str]] = defaultdict(list)
    for obj in objects:
        typ = legacy.object_type(obj)
        if typ in {"domain", "inetnum", "inet6num"} and obj.get("org"):
            key = legacy.primary_value(obj, typ)
            for org_handle in obj["org"]:
                resources_by_org[org_handle].append(f"{typ}: {key}")

    field_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, str]] = []
    object_rows: dict[tuple[str, str], dict[str, str]] = {}
    warnings: list[str] = []

    for index, item in enumerate(legacy_conflicts, start=1):
        cid = f"F{index:03d}"
        target = (item["object_type"], item["object_key"])
        field = item["field"]
        obj = merger.by_target.get(target)
        if not obj:
            warnings.append(f"{cid}: baseline object not found: {target}")
            continue
        evidence = contributions.get((target[0], target[1], field), [])
        candidates = list(dict.fromkeys(source_display(row, field) for row in evidence if source_display(row, field)))
        raw_candidates = list(dict.fromkeys(row["candidate"] for row in evidence if row.get("candidate")))
        original_values = obj.get(field, [])
        original = "\n".join(original_values)
        if field in {"admin-c", "tech-c"} and original_values:
            expanded = [f"{handle}\n{contact_text(contact_objects.get(handle))}" for handle in original_values]
            original = "\n\n".join(expanded)
        reason, action = explanation(field, original, raw_candidates)
        source_refs = list(dict.fromkeys(f"{row['source_sheet']}!{row['source_row']}" for row in evidence))
        units = list(dict.fromkeys(row.get("unit_name", "") for row in evidence if row.get("unit_name")))
        resources = resources_by_org.get(target[1], [])
        field_rows.append({
            "conflict_id": cid,
            "legacy_report_row": item["legacy_row"],
            "object_key": target[1],
            "object_type": target[0],
            "org_name": "\n".join(obj.get("org-name", [])),
            "c_org_name": "\n".join(obj.get("c-org-name", [])),
            "field": field,
            "whois_original": original,
            "filing_candidates": "\n".join(raw_candidates),
            "candidate_count": len(raw_candidates),
            "source_record_count": len(evidence),
            "source_rows": "、".join(source_refs),
            "filing_units": "\n".join(units),
            "associated_resource_count": len(resources),
            "associated_resources": "\n".join(resources),
            "reason": reason,
            "recommended_action": action,
        })
        for evidence_index, source in enumerate(evidence, start=1):
            source_rows.append({"conflict_id": cid, "evidence_id": f"{cid}-{evidence_index:03d}", **source})
        object_rows[target] = {
            "object_key": target[1],
            "object_type": target[0],
            "org_name": "\n".join(obj.get("org-name", [])),
            "c_org_name": "\n".join(obj.get("c-org-name", [])),
            "org_type": "\n".join(obj.get("org-type", [])),
            "cert_type": "\n".join(obj.get("cert-type", [])),
            "cert_no": "\n".join(obj.get("cert-no", [])),
            "admin_c": "\n".join(obj.get("admin-c", [])),
            "tech_c": "\n".join(obj.get("tech-c", [])),
            "c_address": "\n".join(obj.get("c-address", [])),
            "associated_resource_count": str(len(resources)),
            "associated_resources": "\n".join(resources),
            "raw_object": raw_object(obj),
        }

    missing_evidence = [row["conflict_id"] for row in field_rows if row["source_record_count"] == 0]
    if missing_evidence:
        warnings.append("No source evidence reconstructed for: " + ", ".join(missing_evidence))

    by_field: dict[str, int] = defaultdict(int)
    for row in field_rows:
        by_field[row["field"]] += 1
    return {
        "metadata": {
            "baseline": args.baseline.name,
            "icp_workbook": args.icp_workbook.name,
            "ip_workbook": args.ip_workbook.name,
            "original_report": args.original_report.name,
            "conflict_objects": len({(row['object_type'], row['object_key']) for row in field_rows}),
            "conflict_fields": len(field_rows),
            "source_records": len(source_rows),
            "counts_by_field": dict(sorted(by_field.items())),
            "warnings": warnings,
        },
        "conflicts": field_rows,
        "source_records": source_rows,
        "whois_objects": list(object_rows.values()),
        "legacy_report_rows": read_legacy_report_rows(args.original_report),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--legacy-script", type=Path, required=True)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--icp-workbook", type=Path, required=True)
    parser.add_argument("--ip-workbook", type=Path, required=True)
    parser.add_argument("--original-report", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    return parser.parse_args()


if __name__ == "__main__":
    parsed = parse_args()
    payload = build(parsed)
    parsed.output_json.parent.mkdir(parents=True, exist_ok=True)
    parsed.output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload["metadata"], ensure_ascii=False, indent=2))
